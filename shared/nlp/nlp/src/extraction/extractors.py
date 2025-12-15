import abc
import json
import uuid
from typing import List, Dict, Optional

class BaseExtractor(abc.ABC):
    @abc.abstractmethod
    def extract(self, file_path: str, source_doc_id: str) -> List[Dict]:
        """
        Extract content from file.
        Returns list of dicts corresponding to TextChunk schema (excluding DB-specific IDs).
        """
        pass

    def create_chunk(self, 
                     source_doc_id: str, 
                     source_type: str, 
                     raw_text: str, 
                     location_metadata: Dict) -> Dict:
        """Helper to format a chunk dict."""
        return {
            "chunk_id": str(uuid.uuid4()),
            "source_doc_id": source_doc_id,
            "source_type": source_type,
            "location_metadata": json.dumps(location_metadata),
            "raw_text": raw_text,
            "normalized_text": " ".join(raw_text.split()), # Basic normalization
            "is_in_domain": None,
            "link_metadata": "[]",
            "canonical_id": None, # Will be set to chunk_id by Manager if needed
            "improved_text": None
        }

class PDFExtractor(BaseExtractor):
    def extract(self, file_path: str, source_doc_id: str) -> List[Dict]:
        chunks = []
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    chunks.append(self.create_chunk(
                        source_doc_id=source_doc_id,
                        source_type="PDF",
                        raw_text=text,
                        location_metadata={"page": i + 1}
                    ))
        except Exception as e:
            # In a real system, we might log this better or re-raise
            print(f"Error reading PDF {file_path}: {e}")
            raise e
        return chunks

class AudioTranscriber(BaseExtractor):
    def __init__(self, cost_per_min: float = 0.006):
        self.cost_per_min = cost_per_min
        self.accumulated_minutes = 0.0
        self.accumulated_cost = 0.0

    def estimate_cost(self, file_path: str) -> float:
        # Heuristic: file size / bitrate or just mock duration?
        # We can't easily get duration without a library like pydub/ffmpeg
        # For this exercise, we'll try to guess duration or use a placeholder
        # Let's assume 1MB ~= 1 minute for mp3 as a rough heuristic if we can't read it
        size_mb = os.path.getsize(file_path) / (1024 * 1024)
        duration_min = size_mb * 1.0 # Very rough guess
        return duration_min * self.cost_per_min

    def extract(self, file_path: str, source_doc_id: str) -> List[Dict]:
        # Simulating transcription since we can't run Whisper
        # In a real deployment, this would call OpenAI API or local model
        
        # 1. Update stats
        import os
        size_mb = os.path.getsize(file_path) / (1024 * 1024)
        duration_min = size_mb # Mock duration
        self.accumulated_minutes += duration_min
        self.accumulated_cost += duration_min * self.cost_per_min

        # 2. Return dummy transcript
        return [self.create_chunk(
            source_doc_id=source_doc_id,
            source_type="AUDIO",
            raw_text="[TRANSCRIPTION STUB] Content of audio file...",
            location_metadata={"timestamp_start": "00:00:00", "timestamp_end": "00:01:00"}
        )]

class ImageOCR(BaseExtractor):
    def extract(self, file_path: str, source_doc_id: str) -> List[Dict]:
        # Stub for OCR
        return [self.create_chunk(
            source_doc_id=source_doc_id,
            source_type="IMAGE",
            raw_text="[OCR STUB] Text from image...",
            location_metadata={"region": "full"}
        )]

class PPTXExtractor(BaseExtractor):
    def extract(self, file_path: str, source_doc_id: str) -> List[Dict]:
        chunks = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for i, slide in enumerate(prs.slides):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        texts.append(shape.text)
                
                full_text = "\n".join(texts)
                if full_text and full_text.strip():
                     chunks.append(self.create_chunk(
                        source_doc_id=source_doc_id,
                        source_type="PPT",
                        raw_text=full_text,
                        location_metadata={"slide": i + 1}
                    ))
        except Exception as e:
            print(f"Error reading PPT {file_path}: {e}")
            raise e
        return chunks

def get_extractor(source_type: str) -> Optional[BaseExtractor]:
    if source_type == 'PDF':
        return PDFExtractor()
    elif source_type == 'PPT':
        return PPTXExtractor()
    elif source_type == 'AUDIO':
        return AudioTranscriber()
    elif source_type == 'IMAGE':
        return ImageOCR()
    elif source_type == 'TEXT':
         # Simple text reader
         class TextRead(BaseExtractor):
             def extract(self, file_path, sid):
                 with open(file_path, 'r') as f:
                     return [self.create_chunk(sid, 'TEXT', f.read(), {'section': 'full'})]
         return TextRead()
    return None
